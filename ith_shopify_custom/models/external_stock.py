from odoo import api, fields, models
from odoo.exceptions import ValidationError


class ITHExternalStock(models.Model):
    _name = "ith.external.stock"
    _description = "Supplier Stock Availability"
    _rec_name = "sku"
    _order = "sku"

    sku = fields.Char(required=True, index=True)
    quantity = fields.Float(
        string="Supplier Available Quantity",
        required=True,
        default=0.0,
    )
    snapshot_version = fields.Integer(
        string="Stock Snapshot Version",
        default=0,
        readonly=True,
    )

    updated_at = fields.Datetime(
        string="Last Updated",
        default=fields.Datetime.now,
        readonly=True,
    )

    _unique_sku = models.Constraint(
        "UNIQUE(sku)",
        "Only one supplier availability record is allowed per SKU.",
    )

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if "sku" in vals:
                vals["sku"] = (vals["sku"] or "").strip()
            vals["updated_at"] = fields.Datetime.now()
        return super().create(vals_list)

    def write(self, vals):
        vals = dict(vals)
        if "sku" in vals:
            vals["sku"] = (vals["sku"] or "").strip()
        vals["updated_at"] = fields.Datetime.now()
        return super().write(vals)

    @api.constrains("sku")
    def _check_sku(self):
        for record in self:
            if not record.sku:
                raise ValidationError(
                    self.env._("Supplier stock SKU cannot be empty.")
                )

    @api.model
    def replace_supplier_snapshot(self, rows, *, zero_missing=True):
        """Replace the complete external supplier stock snapshot.

        rows: iterable of {"sku": str, "quantity": number}.
        Only ith.external.stock is modified; stock.quant is untouched.
        """
        if not isinstance(rows, (list, tuple)) or not rows:
            raise ValidationError(
                self.env._("A full supplier stock file cannot be empty.")
            )

        incoming = {}

        for row in rows:
            if not isinstance(row, dict):
                raise ValidationError(
                    self.env._("Invalid supplier stock row.")
                )

            sku = str(row.get("sku") or "").strip()
            if not sku:
                raise ValidationError(
                    self.env._("Supplier stock row has an empty SKU.")
                )

            if sku in incoming:
                raise ValidationError(
                    self.env._("Duplicate SKU in supplier file: %s", sku)
                )

            try:
                quantity = float(row["quantity"])
            except (KeyError, TypeError, ValueError):
                raise ValidationError(
                    self.env._("Invalid quantity for SKU %s", sku)
                )

            if not 0 <= quantity < float("inf"):
                raise ValidationError(
                    self.env._("Invalid quantity for SKU %s", sku)
                )

            incoming[sku] = quantity

        with self.env.cr.savepoint():
            self.env.cr.execute(
                "SELECT pg_advisory_xact_lock(hashtext(%s))",
                ("ith.external.stock.snapshot",),
            )

            existing = {r.sku: r for r in self.sudo().search([])}
            zeroed = 0
            next_version = max((r.snapshot_version for r in existing.values()), default=0) + 1

            for sku, quantity in incoming.items():
                if sku in existing:
                    existing[sku].write({
                        "quantity": quantity,
                        "snapshot_version": next_version,
                    })
                else:
                    self.sudo().create({
                        "sku": sku,
                        "quantity": quantity,
                        "snapshot_version": next_version,
                    })

            if zero_missing:
                for sku, record in existing.items():
                    if sku not in incoming:
                        if record.quantity:
                            zeroed += 1
                        record.write({
                            "quantity": 0,
                            "snapshot_version": next_version,
                        })

        return {
            "imported": len(incoming),
            "zeroed": zeroed,
        }


    @api.model
    @api.model
    def reconcile_order_quantity(self, sku, quantity, source, reference, source_line_key):
        """Reconcile order availability once, without replay after an Excel reset."""
        from odoo.exceptions import UserError

        sku = (sku or "").strip()
        if source not in ("shopify", "amazon", "odoo"):
            raise UserError(self.env._("Unsupported order source."))
        if not sku or not reference or not source_line_key:
            raise UserError(self.env._("Missing order reference."))

        try:
            quantity = float(quantity)
        except (TypeError, ValueError):
            raise UserError(self.env._("Invalid order quantity."))
        if not 0 <= quantity < float("inf") or not quantity.is_integer():
            raise UserError(self.env._("Order quantity must be a nonnegative integer."))

        with self.env.cr.savepoint():
            self.env.cr.execute(
                "SELECT pg_advisory_xact_lock(hashtext(%s))",
                ("ith.external.stock.snapshot",),
            )
            stock = self.sudo().search([("sku", "=", sku)], limit=1)
            if not stock:
                raise UserError(self.env._("External SKU %s not found.", sku))

            Movement = self.env["ith.external.stock.movement"].sudo()
            movement = Movement.search([
                ("source", "=", source),
                ("source_line_key", "=", source_line_key),
            ], limit=1)

            if movement and movement.stock_id.id != stock.id:
                raise UserError(self.env._("Order line SKU changed; manual review required."))

            if movement and movement.snapshot_version != stock.snapshot_version:
                # Excel supersedes all availability effects of this old line.
                # Keep the old snapshot version to prevent future replay.
                movement.write({
                    "order_quantity": quantity,
                    "shortage_quantity": 0.0,
                    "quantity_delta": 0.0,
                })
                return {
                    "applied_delta": 0.0,
                    "quantity": stock.quantity,
                    "shortage": 0.0,
                }

            previous = movement.order_quantity if movement else 0.0
            previous_shortage = movement.shortage_quantity if movement else 0.0
            previously_deducted = max(0.0, previous - previous_shortage)

            # Return only the previously deducted quantity, not the shortfall.
            available = stock.quantity + previously_deducted
            deducted = min(quantity, available)
            shortage = quantity - deducted
            new_stock = max(0.0, available - deducted)

            values = {
                "order_quantity": quantity,
                "shortage_quantity": shortage,
                "quantity_delta": -deducted,
            }

            if movement:
                movement.write(values)
            else:
                values.update({
                    "stock_id": stock.id,
                    "source": source,
                    "reference": reference,
                    "source_line_key": source_line_key,
                    "snapshot_version": stock.snapshot_version,
                    "last_applied_snapshot_version": stock.snapshot_version,
                    "applied": True,
                })
                Movement.create(values)

            if new_stock != stock.quantity:
                stock.write({"quantity": new_stock})

            return {
                "applied_delta": deducted - previously_deducted,
                "quantity": new_stock,
                "shortage": shortage,
            }

    @api.model
    def import_supplier_xlsx(self, content, *, zero_missing=False):
        """Import a complete supplier availability Excel snapshot."""
        from io import BytesIO
        import re
        from openpyxl import load_workbook

        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
        try:
            sheet = workbook["Order"]
            rows = sheet.iter_rows(values_only=True)
            headers = [
                str(value or "").strip()
                for value in next(rows)
            ]

            required = ("Internal Reference", "Availability")
            if not all(name in headers for name in required):
                raise ValidationError(
                    self.env._("Required Excel columns are missing.")
                )

            sku_col = headers.index("Internal Reference")
            qty_col = headers.index("Availability")
            incoming = []

            for line_no, row in enumerate(rows, start=2):
                if not any(value is not None for value in row):
                    continue

                sku = str(row[sku_col] or "").strip()
                value = str(row[qty_col] or "").strip()

                if not sku or not value:
                    raise ValidationError(
                        self.env._("Missing SKU or availability at row %s", line_no)
                    )

                if value.lower() == "out of stock":
                    quantity = 0
                elif re.fullmatch(
                    r"over\s+15(?:\.0+)?\s+units available",
                    value,
                    re.IGNORECASE,
                ):
                    quantity = 15
                else:
                    match = re.fullmatch(
                        r"(\d+)\s+units? available",
                        value,
                        re.IGNORECASE,
                    )
                    if not match:
                        raise ValidationError(
                            self.env._(
                                "Unknown availability '%(value)s' at row %(row)s",
                                value=value,
                                row=line_no,
                            )
                        )
                    quantity = int(match.group(1))

                incoming.append({"sku": sku, "quantity": quantity})

            if zero_missing and len(incoming) < 1000:
                raise ValidationError(
                    self.env._(
                        "Full supplier snapshot has fewer than 1000 rows. "
                        "Import refused to prevent accidental stock zeroing."
                    )
                )

            return self.replace_supplier_snapshot(incoming, zero_missing=zero_missing)
        finally:
            workbook.close()

    @api.constrains("quantity")
    def _check_quantity(self):
        for record in self:
            if record.quantity < 0:
                raise ValidationError(
                    self.env._("Supplier availability cannot be negative.")
                )
